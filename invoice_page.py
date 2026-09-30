                                                                     
from customtkinter import *
                                                           
from PIL import Image
                                                               
from tkinter import messagebox
                                                                     
from Tuitionfee_DB import tuitionDB
                                                                     
import os
                                                                     
import re
                                                                             
import subprocess
                                                                      
import sys
                                                                
from datetime import datetime
                                                            
from docx import Document
                                                                      
from docx.enum.text import WD_ALIGN_PARAGRAPH
                                                                       
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
                                                                   
from docx.shared import Inches, Pt
                                                               
from docxtpl import DocxTemplate



                                                                         
class InvoicePage(CTkFrame):
                                                    
    def __init__(self, parent,db):
                                                                  
        super().__init__(parent, fg_color="transparent")
                                                                               
        self.db = db
                                                                               
        self.selected_id = None
                                                                               
        self.selected_card = None

       
                                                                               
        self.columnconfigure(0, weight=2)
                                                                               
        self.columnconfigure(1, weight=3)
                                                                               
        self.rowconfigure(0, weight=0)              
                                                                               
        self.rowconfigure(1, weight=1)               

                 
                                                        
        cards_row = CTkFrame(self, fg_color="transparent")
                                                                      
        cards_row.grid(row=0, column=0, padx=15, pady=(15, 5), sticky="ew")

                                                            
        for i in range(5):
                                                                      
            cards_row.columnconfigure(i, weight=1) 

                                                                               
        self.stat_cards = {}
                                                      
        cards_data = [
                                                                      
            ("#2563eb",
                                                                       
             "multiple-users-silhouette.png",
                                                                       
             "Total Students",
                                                                       
             "All Registered Students",self.show_all_students),

                                                                      
            ("#dc2626",
                                                                       
             "user.png",
                                                                       
             "Unpaid Students",
                                                                       
             "Who haven't paid any fee", self.show_unpaid_students),

                                                                      
            ("#16a34a",
                                                                       
             "accept.png", 
                                                                       
             "Paid Students", 
                                                                       
             "Fee fully paid",self.show_paid_students),

                                                                      
            ("#f59e0b",
                                                                       
             None,
                                                                       
             "All Recent Invoices",
                                                                       
             "Latest fee receipts", self.refresh_recent_invoices),
                                                                            
        ]

                                                            
        for i, (color, icon, title, sub, command) in enumerate(cards_data):
                                                        
            card = self.create_stat_card(
                                                                          
                cards_row, color, icon, "0", title, sub, command
                                                                                
            )
                                                                                   
            self.stat_cards[title] = card
                                                                          
            card.grid(row=0, column=i, padx=5, pady=5, sticky="nsew")

                                                        
        left_panel = CTkFrame(self, fg_color="#0b1a2e", corner_radius=10,
                                                                    
                        border_width=1, border_color="#1f3a5f")
                                                                      
        left_panel.grid(row=1, column=0, padx=(15,7), pady=(5,15), sticky="nsew")
                                                                  
        left_panel.columnconfigure(0, weight=1)
                                                                  
        left_panel.rowconfigure(1, weight=1)

                                                        
        recent_header = CTkFrame(left_panel, fg_color="transparent")
                                                                      
        recent_header.grid(row=0, column=0, sticky="ew", padx=15, pady=(15, 10))
                                                                  
        recent_header.columnconfigure(0, weight=1)

                                                        
        CTkLabel(recent_header, text="📄, Recent Fee Collections", font=("Arial Black", 13),
                                                                          
            text_color="#eeeee4").grid(row=0, column=0, sticky="w")

                                                        
        CTkButton(
                                                                      
            recent_header,
                                                        
            text="Download",
                                                        
            width=90,
                                                        
            height=28,
                                                        
            fg_color="#12263f",
                                                        
            hover_color="#1f3a5f",
                                                        
            text_color="white",
                                                        
            font=("Arial Black", 10),
                                                        
            command=self.download_selected_invoice,
                                                                      
        ).grid(row=0, column=1, sticky="e", padx=(8, 0))

                                                        
        CTkButton(
                                                                      
            recent_header,
                                                        
            text="Edit",
                                                        
            width=70,
                                                        
            height=28,
                                                        
            fg_color="#12263f",
                                                        
            hover_color="#1f3a5f",
                                                        
            text_color="white",
                                                        
            font=("Arial Black", 10),
                                                        
            command=lambda: self.handle_invoice_action("Edit"),
                                                                      
        ).grid(row=0, column=2, sticky="e", padx=(8, 0))

                                                        
        CTkButton(
                                                                      
            recent_header,
                                                        
            text="Delete",
                                                        
            width=70,
                                                        
            height=28,
                                                        
            fg_color="#dc2626",
                                                        
            hover_color="#b91c1c",
                                                        
            text_color="white",
                                                        
            font=("Arial Black", 10),
                                                        
            command=lambda: self.handle_invoice_action("Delete"),
                                                                      
        ).grid(row=0, column=3, sticky="e", padx=(8, 0))

                                          
                                                                               
        self.entries_frame = CTkScrollableFrame(left_panel, fg_color="transparent")
                                                                               
        self.entries_frame.grid(row=1, column=0, sticky="nsew", padx=10, pady=(0,15))
                                                                               
        self.entries_frame.columnconfigure(0, weight=1)

                                                                               
        self.refresh_stat_cards()
                                                                               
        self.show_all_students()
                                                                                
        if "Total Students" in self.stat_cards:
                                                                                   
            self.selected_card = self.stat_cards["Total Students"]
                                                                                   
            self.selected_card.configure(border_color="#38bdf8")


                                                                               
        self.refresh_recent_invoices()


                                                        
        right_panel = CTkScrollableFrame(self, fg_color="#0b1a2e", corner_radius=10,
                                                                    
                        border_width=1, border_color="#1f3a5f", width=500)
                                                                      
        right_panel.grid(row=0, column=1, rowspan=2, padx=(7,15), pady=(15,15), sticky="nsew")

                                                                  
        right_panel.columnconfigure(0, weight=1)



                                                        
        details_frame = CTkFrame(right_panel, fg_color="#0b1a2e", corner_radius=8,
                                                                             
                                border_width=1, border_color="#1f3a5f")
                                                                      
        details_frame.grid(row=2, column=0, columnspan=3, sticky="ew", padx=15, pady=(0,8))
                                                                               
        self.detail_labels = {}

                                                      
        left_fields = ["ID", "Name", "Class"]
                                                      
        right_fields = ["Phone_No", "Monthly Fee", "Status"]

                                                            
        for i, field in enumerate(left_fields):
                                                            
            CTkLabel(details_frame, text=f"{field}:", font=("Arial Black", 12),
                                                                                  
                    text_color="#eeeee4").grid(row=i, column=0, sticky="w", padx=(20,5), pady=3)
                                                            
            val_label = CTkLabel(details_frame, text="-", font=("Arial",11), text_color="#eeeee4")
                                                                          
            val_label.grid(row=i, column=1, sticky="w", padx=10, pady=3)
                                                                                   
            self.detail_labels[field] = val_label

                                                            
        for i, field in enumerate(right_fields):
                                                            
            CTkLabel(details_frame, text=f"{field}:", font=("Arial Black", 12),
                                                                                  
                    text_color="#eeeee4").grid(row=i, column=2, sticky="w", padx=(50,40), pady=3)
                                                            
            val_label = CTkLabel(details_frame, text="-", font=("Arial",11), text_color="#eeeee4")
                                                                          
            val_label.grid(row=i, column=3, sticky="w",columnspan=2, padx=10, pady=3)
                                                                                   
            self.detail_labels[field] = val_label





        
               
                                                        
        CTkLabel(right_panel, text="Create Receipt", font=("Arial Black", 13),
                                                                              
                text_color="#eeeee4").grid(row=0, column=0, columnspan=3, sticky="w", padx=15, pady=(15,6))
        
                        
                                                        
        search_frame = CTkFrame(right_panel, fg_color="transparent")
                                                                      
        search_frame.grid(row=1, column=0, columnspan=3, sticky="ew", padx=15, pady=(0,8))
                                                                  
        search_frame.columnconfigure(1, weight=1)
        
                                                                               
        self.search_entry = CTkEntry(
                                                                      
            search_frame, placeholder_text="🔎︎ Search...",
                                                        
            fg_color="#12263f", border_color="#1f3a5f", text_color="white"
                                                                            
        )
                                                                               
        self.search_entry.grid(row=0, column=0, sticky="ew", padx=(0,10))

                                                                               
        self.search_option = CTkOptionMenu(search_frame, values=["ID", "Name", "Class"],
                                                        
            fg_color="#12263f",
                                                        
            button_color="#1f3a5f",
                                                        
            text_color="white",
                                                                            
        )
                                                                               
        self.search_option.grid(row=0, column=1, sticky="ew", padx=(0,10))

                                                        
        CTkButton(search_frame,
                                                        
            text="Search",
                                                        
            fg_color="#f78222",
                                                        
            hover_color="#d4691a",
                                                        
            text_color="black",
                                                        
            font=("Arial Black", 12),
                                                                          
            command=self.search_student).grid(row=0, column=2, sticky="ew")

                                                                               
        self.selected_invoice = None
                                                                               
        self.selected_invoice_id = None
                                                                     
                                                                                             
                                                  
                                                  
                                                   


                                                        
        method_frame = CTkFrame(right_panel, fg_color="transparent")
                                                                      
        method_frame.grid(row=3, column=0, sticky="w", padx=15, pady=(15,10))

                                                        
        CTkLabel(method_frame, text="Payment Method", font=("Arial Black", 12),
                                                                              
                text_color="#eeeee4").grid(row=0, column=0, sticky="w")

                                                                               
        self.payment_method=CTkOptionMenu(method_frame,
                                                                
                    font=("Arial Black", 9)
                                                                  
                    ,values=["Cash","UPI","Cards","Other"],
                                                                
                    text_color="#eeeee4",
                                                                
                    fg_color="#12263f",
                                                                
                    button_color="#1f3a5f",
                                                                                        
                    )
                                                                               
        self.payment_method.grid(row=1, column=0, sticky="w", pady=(10,0))

                                                        
        date_col_frame = CTkFrame(right_panel, fg_color="transparent")
                                                                      
        date_col_frame.grid(row=3, column=1, sticky="w", padx=15, pady=(6))

                                                        
        CTkLabel(date_col_frame, text="Payment Date", font=("Arial Black", 12),
                                                                                      
                        text_color="#eeeee4").grid(row=0, column=0, sticky="w")

                                                        
        date_frame = CTkFrame(date_col_frame, fg_color="transparent")
                                                                      
        date_frame.grid(row=1, column=0, sticky="w", pady=(6))

                                                                               
        self.date_combo = CTkOptionMenu(date_frame, width=70, height=20,
                                                          
            values=['Days','01','02','03','04','05','06','07','08','09','10',
                                                                              
                    '11','12','13','14','15','16','17','18','19','20',
                                                                              
                    '21','22','23','24','25','26','27','28','29','30','31'],
                                                        
            font=("Arial Black", 9),
                                                        
            text_color="#eeeee4",
                                                        
            fg_color="#12263f",
                                                        
            button_color="#1f3a5f",
                                                                                
            )
                                                                               
        self.date_combo.grid(row=0, column=0, padx=2)

                                                                               
        self.month_combo = CTkOptionMenu(date_frame, width=70, height=20,
                                                                  
                    values=['Month','01','02','03','04','05','06','07','08','09','10','11','12'],
                                                                
                    font=("Arial Black", 9),
                                                                
                    text_color="#eeeee4",
                                                                
                    fg_color="#12263f",
                                                                
                    button_color="#1f3a5f",
                                                                                        
                    )
                                                                               
        self.month_combo.grid(row=0, column=1, padx=2)

                                                                               
        self.year_combo = CTkOptionMenu(date_frame, width=70, height=20,
                                                                  
                    values=['Years','2026','2027','2028','2029','2030','2031'],
                                                                
                    font=("Arial Black", 9),
                                                                
                    text_color="#eeeee4",
                                                                
                    fg_color="#12263f",
                                                                
                    button_color="#1f3a5f",
                                                                                        
                    )
                                                                               
        self.year_combo.grid(row=0, column=2, padx=2)
        
                                                        
        note_frame = CTkFrame(right_panel, fg_color="transparent")
                                                                      
        note_frame.grid(row=4, column=0, columnspan=1,sticky="w", padx=15, pady=(15,10))

                                                        
        CTkLabel(note_frame,text="Note(Option)", font=("Arial Black", 12),
                                                                              
                text_color="#eeeee4").grid(row=0, column=0,pady=6)
        
                                                                               
        self.note_textbox = CTkTextbox(note_frame, width=350, height= 25,
                                                            
                font=("Arial Black", 9),
                                                            
                text_color="#000000",
                                                            
                corner_radius = 2)
                                                                               
        self.note_textbox.grid(row=0, column=1,padx= 15, pady=6)

                                                        
        submit_frame = CTkFrame(right_panel, fg_color="transparent")
                                                                      
        submit_frame.grid(row=5, column=0, columnspan=1,sticky="w", padx=15, pady=(6))

                                                                               
        self.submit_button = CTkButton(submit_frame, width=70, height=20,
                                                                                
                                    corner_radius=10, bg_color="transparent",
                                                                                
                                    text="Submit", fg_color="#0907f2",
                                                                                
                                    font=("Arial Black", 15),
                                                                                
                                    text_color="#eeeee4",
                                                                                
                                    command=self.submit_payment)
                                                                               
        self.submit_button.grid(row=0,column =0,sticky="ew", padx=25, pady=6)

                                                        
        history_panel = CTkFrame(
                                                                      
            right_panel,
                                                        
            fg_color="#0b1a2e",
                                                        
            corner_radius=8,
                                                        
            border_width=1,
                                                        
            border_color="#1f3a5f",
                                                                            
        )
                                                                      
        history_panel.grid(row=6, column=0, columnspan=3, sticky="nsew", padx=15, pady=(8, 15))
                                                                  
        history_panel.columnconfigure(0, weight=1)

                                                        
        history_header = CTkFrame(history_panel, fg_color="transparent")
                                                                      
        history_header.grid(row=0, column=0, sticky="ew", padx=12, pady=(10, 4))
                                                                  
        history_header.columnconfigure(0, weight=1)

                                                        
        CTkLabel(
                                                                      
            history_header,
                                                        
            text="All Recent Invoices",
                                                        
            font=("Arial Black", 13),
                                                        
            text_color="#eeeee4",
                                                                      
        ).grid(row=0, column=0, sticky="w")

                                                        
        CTkButton(
                                                                      
            history_header,
                                                        
            text="Refresh",
                                                        
            width=80,
                                                        
            height=24,
                                                        
            fg_color="#12263f",
                                                        
            hover_color="#1f3a5f",
                                                        
            text_color="#eeeee4",
                                                        
            font=("Arial Black", 10),
                                                        
            command=self.refresh_recent_invoice_details,
                                                                      
        ).grid(row=0, column=1, sticky="e")

                                                                               
        self.invoice_history_frame = CTkScrollableFrame(
                                                                      
            history_panel,
                                                        
            fg_color="transparent",
                                                        
            height=170,
                                                                            
        )
                                                                               
        self.invoice_history_frame.grid(row=1, column=0, sticky="nsew", padx=8, pady=(0, 10))
                                                                               
        self.invoice_history_frame.columnconfigure(0, weight=1)
                                                                               
        self.refresh_recent_invoice_details()






                           




                                                            
    def create_stat_card(self, parent, icon_color, icon_path, number, title, subtitle, command=None):
                                                        
        card = CTkFrame(
                                                                      
            parent,
                                                        
            fg_color="#0b1a2e",
                                                        
            corner_radius=10,
                                                        
            border_width=1,
                                                        
            border_color="#1f3a5f",
                                                        
            width=180,
                                                        
            height=110,
                                                                            
        )
                                                                  
        card.grid_columnconfigure(0, weight=1)

                                                        
        icon_badge = CTkFrame(card, fg_color=icon_color, corner_radius=8, width=32, height=32)
                                                                      
        icon_badge.grid(row=0, column=0, sticky="w", padx=10, pady=(10, 0))

                                                                                
        if icon_path:
                                                                       
            try:
                                                                
                icon_img = CTkImage(light_image=Image.open(icon_path), size=(18, 18))
                                                                       
                CTkLabel(icon_badge, text="", image=icon_img).place(relx=0.5, rely=0.5, anchor="center")
                                                            
            except Exception:
                                                                       
                CTkLabel(icon_badge, text="â—", text_color="white", font=("Arial Black", 12)).place(relx=0.5, rely=0.5, anchor="center")

                                                                               
        self.number_label = CTkLabel(card, text=str(number), font=("Arial Black", 20), text_color="#eeeee4")
                                                                               
        self.number_label.grid(row=1, column=0, sticky="w", padx=10)
                                                                      
        CTkLabel(card, text=title, font=("Arial", 12, "bold"), text_color="#eeeee4").grid(row=2, column=0, sticky="w", padx=10)
                                                                      
        CTkLabel(card, text=subtitle, font=("Arial", 9), text_color="#7d8ba1").grid(row=3, column=0, sticky="w", padx=10, pady=(0, 10))

                                                                  
        card.bind("<Button-1>", lambda event: self.select_stat_card(card, command))
                                                            
        for widget in card.winfo_children():
                                                                      
            widget.bind("<Button-1>", lambda event: self.select_stat_card(card, command))

                                                    
        card.count_label = self.number_label
                                                      
        return card

                                                            
    def select_stat_card(self, card, command):
                                                                                
        if command is not None:
                                                                      
            command()

                                                                                
        if hasattr(self, "selected_card") and self.selected_card is not None and self.selected_card != card:
                                                                                   
            self.selected_card.configure(border_color="#1f3a5f")

                                                                  
        card.configure(border_color="#38bdf8")
                                                                               
        self.selected_card = card

                                                              
    def refresh_stat_cards(self):
                                                    
        students = self.db.get_all_students() if hasattr(self.db, 'get_all_students') else []
                                                    
        total = len(students)
                                                    
        paid = sum(1 for s in students if s[5] == 'P')
                                                    
        unpaid = sum(1 for s in students if s[5] == 'UN')

                                                            
        for title, card in self.stat_cards.items():
                                                                                    
            if title == "Total Students":
                                                                          
                card.count_label.configure(text=str(total))
                                                                                       
            elif title == "Paid Students":
                                                                          
                card.count_label.configure(text=str(paid))
                                                                                       
            elif title == "Unpaid Students":
                                                                          
                card.count_label.configure(text=str(unpaid))
                                                                                       
            elif title == "All Recent Invoices":
                                                                                       
                self.update_recent_invoice_count()

                                                                       
    def update_recent_invoice_count(self):
                                                    
        card = self.stat_cards.get("All Recent Invoices") if hasattr(self, "stat_cards") else None
                                                                                
        if not card:
                                                          
            return

                                                                                
        if hasattr(self.db, "count_invoices"):
                                                        
            count = self.db.count_invoices()
                                                                  
        else:
                                                        
            count = len(self.db.get_recent_invoices(100000))
                                                                  
        card.count_label.configure(text=str(count))

                                                          
    def search_student(self):
                                                    
        query = self.search_entry.get().strip()
                                                    
        filter = self.search_option.get()
                                                                                
        if not query:
                                                                      
            messagebox.showwarning("search","Enter a value to search")
                                                          
            return
        
                                                            
        column_map = {"ID":"id","Name":"name","Class":"class"}
                                                    
        col = column_map[filter]


                                                    
        row = self.db.search_students(col, query)

                                                                                
        if not row:
                                                             
            messagebox.showinfo("search", "Not Student Found")
                                                                                   
            self.clear_details()
                                                          
            return
        
                                                                                
        if len(row) == 1:

            self.fill_details(row[0])

                                                        
    def fill_details(self,row):
                                                                               
        self.selected_id = row[0]
                                                                               
        self.detail_labels["ID"].configure(text = row[0])
                                                                               
        self.detail_labels["Name"].configure(text = row[1])
                                                                               
        self.detail_labels["Class"].configure(text = row[2])
                                                                               
        self.detail_labels["Phone_No"].configure(text = row[3])
                                                                               
        self.detail_labels["Monthly Fee"].configure(text = f"₹{row[4]}")
                                                                               
        self.detail_labels["Status"].configure(text = "Paid" if row[5] == "P" else "Unpaid")


                                                         
    def clear_details(self):
                                                                               
        self.selected_id = None
                                                            
        for label in self.detail_labels.values():
                                                                      
            label.configure(text="-")

                                                          
    def submit_payment(self):
                                                                                
        if not self.selected_id:
                                                                      
            messagebox.showwarning("Error", "No Student Can't Search ")
                                                          
            return

                                                    
        day = self.date_combo.get()
                                                    
        month = self.month_combo.get()
                                                    
        year = self.year_combo.get()

                                                                                
        if day == "Days" or month == "Month" or year == "Years":
                                                                      
            messagebox.showwarning("Error", "Payment date Not selected")
                                                          
            return

                                                    
        method = self.payment_method.get()
                                                    
        note = self.note_textbox.get("1.0", "end").strip()

                                                    
        payment_date = f"{year}-{month}-{day.zfill(2)}"
                                                    
        month_label = f"{month}-{year}"

                                                                               
        self.db.insert_invoice(self.selected_id, month_label, "P", method, note, payment_date)
                                                         
        messagebox.showinfo("Success", "Payment Record Successfully ")
                                                                               
        self.clear_details()
                                                                               
        self.note_textbox.delete("1.0", "end")
                                                                               
        self.refresh_recent_invoices()
                                                                               
        self.refresh_stat_cards()
            
                                                             
    def show_all_students(self):
                                                    
        students = self.db.get_all_students()
                                                                               
        self.display_students(students)


                                                              
    def show_paid_students(self):
        students = self.db.get_student_by_status("P")
        self.display_students(students)


                                                                
    def show_unpaid_students(self):
                                                    
        students = self.db.get_student_by_status("UN")
                                                                               
        self.display_students(students)


                                                            
    def display_students(self, students):
                                                            
        for widget in self.entries_frame.winfo_children():
                                                                      
            widget.destroy()

                                                            
        for i, student in enumerate(students):
                                                            
            row = CTkFrame(
                                                                                       
                self.entries_frame,
                                                            
                fg_color="#132540",
                                                            
                corner_radius=8
                                                                                
            )
                                                                          
            row.grid(row=i, column=0, sticky="ew", pady=4, padx=2)
                                                                      
            row.columnconfigure(1, weight=1)

                                                        
            student_id = student[0]
                                                        
            name = student[1]
                                                        
            student_class = student[2]
                                                        
            fee = student[4]
                                                        
            status = student[5]
                                                        
            status_text = "Paid" if status == "P" else "Unpaid"

                                                            
            CTkLabel(
                                                                          
                row,
                                                            
                text=f"ID-{student_id}\nClass {student_class}",
                                                            
                font=("Arial", 10),
                                                            
                text_color="#7d8ba1"
                                                                          
            ).grid(row=0, column=0, padx=10, pady=8, sticky="w")

                                                            
            CTkLabel(
                                                                          
                row,
                                                            
                text=name,
                                                            
                font=("Arial", 11, "bold"),
                                                            
                text_color="#eeeee4"
                                                                          
            ).grid(row=0, column=1, padx=10, pady=8, sticky="w")

                                                            
            CTkLabel(
                                                                          
                row,
                                                            
                text=f"₹ {fee}",
                                                            
                font=("Arial", 11, "bold"),
                                                            
                text_color="#eeeee4"
                                                                          
            ).grid(row=0, column=2, padx=10)

                                                            
            CTkLabel(
                                                                          
                row,
                                                            
                text=status_text,
                                                            
                font=("Arial", 10, "bold"),
                                                            
                text_color="#47d246" if status == "P" else "#dc2626"
                                                                          
            ).grid(row=0, column=3, padx=10)

                                                                      
            row.bind("<Button-1>", lambda event, s=student: self.fill_details(s))
                                                                
            for child in row.winfo_children():
                                                                          
                child.bind("<Button-1>", lambda event, s=student: self.fill_details(s))

                                                                   
    def refresh_recent_invoices(self):
                                                                                
        if not hasattr(self, "entries_frame"):
                                                          
            return

                                                            
        for widget in self.entries_frame.winfo_children():
                                                                      
            widget.destroy()

                                                    
        invoices = self.db.get_recent_invoices(10)
                                                                               
        self.update_recent_invoice_count()
                                                                               
        self.selected_invoice = None
                                                                               
        self.selected_invoice_id = None

                                                            
        for i, invoice in enumerate(invoices):
                                                        
            invoice_id = invoice[9] or f"RCP-{invoice[0]:05d}"
                                                        
            name = invoice[2]
                                                        
            student_class = invoice[3]
                                                        
            amount = invoice[5]
                                                        
            payment_date = invoice[8]
                                                        
            payment_method = invoice[10] or "Cash"
                                                        
            payment_state = invoice[7]
                                                        
            status_text = "Completed" if payment_state == "P" else "Pending"
                                                        
            status_color = "#16a34a" if payment_state == "P" else "#f59e0b"

                                                            
            row = CTkFrame(self.entries_frame, fg_color="#132540", corner_radius=8)
                                                                          
            row.grid(row=i, column=0, sticky="ew", pady=4, padx=2)
                                                                      
            row.columnconfigure(1, weight=1)

                                                                          
            CTkLabel(row, text=f"{invoice_id}\n{payment_date}", font=("Arial", 10), text_color="#7d8ba1", justify="left").grid(row=0, column=0, padx=10, pady=8, sticky="w")
                                                                          
            CTkLabel(row, text=f"{name}\nClass {student_class}", font=("Arial", 11, "bold"), text_color="#eeeee4", justify="left").grid(row=0, column=1, padx=10, pady=8, sticky="w")
                                                                          
            CTkLabel(row, text=f"₹{amount}", font=("Arial", 11, "bold"), text_color="#eeeee4").grid(row=0, column=2, padx=10)
                                                                          
            CTkLabel(row, text=payment_method, font=("Arial", 10), text_color="#7d8ba1").grid(row=0, column=3, padx=10)
                                                                          
            CTkLabel(row, text=status_text, font=("Arial", 10, "bold"), text_color=status_color).grid(row=0, column=4, padx=10)

                                                                      
            row.bind("<Button-1>", lambda event, inv=invoice, rw=row: self.select_recent_invoice(inv, rw))
                                                                
            for child in row.winfo_children():
                                                                          
                child.bind("<Button-1>", lambda event, inv=invoice, rw=row: self.select_recent_invoice(inv, rw))

                                                                                
        if not invoices:
                                                                          
            CTkLabel(self.entries_frame, text="No recent invoice found", font=("Arial", 12), text_color="#7d8ba1").grid(row=0, column=0, padx=10, pady=20)

                                                                               
        self.refresh_recent_invoice_details()

                                                                          
    def refresh_recent_invoice_details(self):
                                                                                
        if not hasattr(self, "invoice_history_frame"):
                                                          
            return

                                                            
        for widget in self.invoice_history_frame.winfo_children():
                                                                      
            widget.destroy()

                                                    
        invoices = self.db.get_recent_invoices(25)
                                                                                
        if not invoices:
                                                            
            CTkLabel(
                                                                                       
                self.invoice_history_frame,
                                                            
                text="No recent invoice found",
                                                            
                font=("Arial", 12),
                                                            
                text_color="#7d8ba1",
                                                                          
            ).grid(row=0, column=0, padx=10, pady=20, sticky="w")
                                                          
            return

                                                            
        for i, invoice in enumerate(invoices):
                                                        
            receipt_no = invoice[9] or f"RCP-{invoice[0]:05d}"
                                                        
            student_name = invoice[2]
                                                        
            student_class = invoice[3]
                                                        
            amount = invoice[5]
                                                        
            month = invoice[6]
                                                        
            status = "Paid" if invoice[7] == "P" else "Pending"
                                                        
            payment_date = invoice[8] or "-"
                                                        
            method = invoice[10] or "Cash"
                                                        
            note = invoice[11] or "-"
                                                        
            updated_at = invoice[12] if len(invoice) > 12 else None
                                                        
            edit_summary = invoice[13] if len(invoice) > 13 else None

                                                            
            row = CTkFrame(
                                                                                       
                self.invoice_history_frame,
                                                            
                fg_color="#132540",
                                                            
                corner_radius=8,
                                                            
                border_width=1,
                                                            
                border_color="#1f3a5f",
                                                                                
            )
                                                                          
            row.grid(row=i, column=0, sticky="ew", padx=2, pady=4)
                                                                      
            row.columnconfigure(1, weight=1)

                                                            
            CTkLabel(
                                                                          
                row,
                                                            
                text=f"{receipt_no}\nID-{invoice[1]}",
                                                            
                font=("Arial", 10),
                                                            
                text_color="#38bdf8",
                                                            
                justify="left",
                                                                          
            ).grid(row=0, column=0, rowspan=2, sticky="nw", padx=10, pady=8)

                                                            
            CTkLabel(
                                                                          
                row,
                                                            
                text=f"{student_name} | Class {student_class} | {month} | Rs.{amount}",
                                                            
                font=("Arial Black", 11),
                                                            
                text_color="#eeeee4",
                                                            
                justify="left",
                                                                          
            ).grid(row=0, column=1, sticky="w", padx=8, pady=(8, 2))

                                                            
            CTkLabel(
                                                                          
                row,
                                                            
                text=f"Date: {payment_date}   Method: {method}   Status: {status}   Note: {note}",
                                                            
                font=("Arial", 10),
                                                            
                text_color="#a9b7ca",
                                                            
                justify="left",
                                                            
                wraplength=420,
                                                                          
            ).grid(row=1, column=1, sticky="w", padx=8, pady=(0, 2))

                                                                                    
            if edit_summary:
                                                            
                edit_text = f"Edited: {updated_at or '-'} | {edit_summary}"
                                                            
                edit_color = "#38bdf8"
                                                                      
            else:
                                                            
                edit_text = "Edited: Not edited yet"
                                                            
                edit_color = "#7d8ba1"

                                                            
            CTkLabel(
                                                                          
                row,
                                                            
                text=edit_text,
                                                            
                font=("Arial", 10, "bold"),
                                                            
                text_color=edit_color,
                                                            
                justify="left",
                                                            
                wraplength=510,
                                                                          
            ).grid(row=2, column=0, columnspan=2, sticky="w", padx=10, pady=(0, 8))

                                                                 
    def select_recent_invoice(self, invoice, row, download=False):
                                                                               
        self.selected_invoice = invoice
                                                                               
        self.selected_invoice_id = invoice[0]

                                                            
        for child in self.entries_frame.winfo_children():
                                                                      
            child.configure(border_color="#1f3a5f", border_width=1)

                                                                  
        row.configure(border_color="#38bdf8", border_width=2)

                                                                                
        if download:
                                                                                   
            self.download_invoice(invoice, show_message=False)

                                                                     
    def download_selected_invoice(self):
                                                                                
        if not self.selected_invoice:
                                                                      
            messagebox.showwarning("Warning", "Please select a recent invoice first.")
                                                          
            return
                                                                               
        self.download_invoice(self.selected_invoice, show_message=True)

                                                            
    def download_invoice(self, invoice, show_message=True):
                                                    
        receipt_no = invoice[9] or f"RCP-{invoice[0]:05d}"
                                                    
        student_id = invoice[1]
                                                    
        student_name = invoice[2]
                                                    
        student_class = invoice[3]
                                                    
        phone = invoice[4]
                                                    
        amount = invoice[5]
                                                    
        month = invoice[6]
                                                    
        status = "Paid" if invoice[7] == "P" else "Pending"
                                                    
        payment_date = invoice[8] or "-"
                                                    
        method = invoice[10] or "Cash"
                                                    
        note = invoice[11] or "-"

                                                    
        safe_receipt = re.sub(r"[^A-Za-z0-9_-]+", "_", str(receipt_no))
                                                    
        safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", str(student_name)).strip("_")
                                                    
        downloads_dir = os.path.join(os.path.expanduser("~"), "Downloads", "FeeTrack Receipts")
                                                                  
        os.makedirs(downloads_dir, exist_ok=True)
                                                    
        base_path = os.path.join(downloads_dir, f"{safe_receipt}_{safe_name}")
                                                    
        docx_path = f"{base_path}.docx"

                                                                  
        create_invoice_docx(
                                                                      
            docx_path,
                                                                      
            {
                                                                          
                "receipt_no": receipt_no,
                                                                          
                "student_id": student_id,
                                                                          
                "student_name": student_name,
                                                                          
                "student_class": student_class,
                                                                          
                "phone": phone,
                                                                          
                "amount": amount,
                                                                          
                "month": month,
                                                                          
                "status": status,
                                                                          
                "payment_date": payment_date,
                                                                          
                "method": method,
                                                                          
                "note": note,
                                                                                
            },
                                                                            
        )
                                                    
        file_path = convert_docx_to_pdf(docx_path, downloads_dir)

                                                                                
        if show_message:
                                                             
            messagebox.showinfo("Downloaded", f"Invoice downloaded:\n{file_path}")
                                                      
        return file_path

                                                                 
    def handle_invoice_action(self, choice):
                                                                                
        if not self.selected_invoice:
                                                                      
            messagebox.showwarning("Warning", "Please select a recent invoice first.")
                                                          
            return

                                                                                
        if choice == "Edit":
                                                                                   
            self.edit_invoice_dialog()
                                                                                   
        elif choice == "Delete":
                                                                                   
            self.delete_invoice_dialog()

                                                               
    def edit_invoice_dialog(self):
                                                                                
        if not self.selected_invoice:
                                                          
            return

                                                    
        invoice = self.selected_invoice
                                                    
        dialog = CTkToplevel(self)
                                                                  
        dialog.title("Edit Invoice")
                                                                  
        dialog.geometry("420x650")
                                                                  
        dialog.grab_set()

                                                                      
        CTkLabel(dialog, text="Edit Invoice", font=("Arial Black", 16), text_color="#eeeee4").pack(pady=(20, 10))

                                                                      
        CTkLabel(dialog, text="Student ID:", text_color="#eeeee4").pack(anchor="w", padx=20, pady=(8, 2))
                                                        
        student_id_var = CTkEntry(dialog, width=300)
                                                                  
        student_id_var.insert(0, str(invoice[1]))
                                                                      
        student_id_var.pack(pady=(0, 10))

                                                                      
        CTkLabel(dialog, text="Date:", text_color="#eeeee4").pack(anchor="w", padx=20, pady=(0, 2))
                                                        
        date_entry = CTkEntry(dialog, width=300)
                                                                  
        date_entry.insert(0, str(invoice[8]))
                                                                      
        date_entry.pack(pady=(0, 10))

                                                                      
        CTkLabel(dialog, text="Month:", text_color="#eeeee4").pack(anchor="w", padx=20, pady=(0, 2))
                                                        
        month_entry = CTkEntry(dialog, width=300)
                                                                  
        month_entry.insert(0, invoice[6])
                                                                      
        month_entry.pack(pady=(0, 10))

                                                                      
        CTkLabel(dialog, text="Amount:", text_color="#eeeee4").pack(anchor="w", padx=20, pady=(0, 2))
                                                        
        amount_entry = CTkEntry(dialog, width=300)
                                                                  
        amount_entry.insert(0, str(invoice[5]))
                                                                      
        amount_entry.pack(pady=(0, 10))

                                                                      
        CTkLabel(dialog, text="Status:", text_color="#eeeee4").pack(anchor="w", padx=20, pady=(0, 2))
                                                      
        status_option = CTkOptionMenu(dialog, values=["P", "UN"], width=300)
                                                                  
        status_option.set(invoice[7])
                                                                      
        status_option.pack(pady=(0, 10))

                                                                      
        CTkLabel(dialog, text="Method:", text_color="#eeeee4").pack(anchor="w", padx=20, pady=(0, 2))
                                                      
        method_option = CTkOptionMenu(dialog, values=["Cash", "UPI", "Cards", "Other"], width=300)
                                                                  
        method_option.set(invoice[10] or "Cash")
                                                                      
        method_option.pack(pady=(0, 10))

                                                                      
        CTkLabel(dialog, text="Note:", text_color="#eeeee4").pack(anchor="w", padx=20, pady=(0, 2))
                                                        
        note_box = CTkTextbox(dialog, width=300, height=70)
                                                                  
        note_box.insert("1.0", invoice[11] or "")
                                                                      
        note_box.pack(pady=(0, 10))

                                                            
        def save_changes():
                                                                       
            try:
                                                            
                new_student_id = int(student_id_var.get().strip())
                                                                                        
                if not self.db.get_student(new_student_id):
                                                              
                    messagebox.showerror("Error", "Student ID not found.")
                                                                  
                    return

                                                            
                new_date = date_entry.get().strip()
                                                            
                new_month = month_entry.get().strip()
                                                            
                new_amount = float(amount_entry.get().strip())
                                                            
                new_status = status_option.get()
                                                            
                new_method = method_option.get()
                                                            
                new_note = note_box.get("1.0", "end").strip()

                                                            
                edit_summary = self.db.update_invoice(
                                                                              
                    invoice[0],
                                                                              
                    new_student_id,
                                                                              
                    new_month,
                                                                              
                    new_amount,
                                                                              
                    new_status,
                                                                              
                    new_date,
                                                                              
                    new_method,
                                                                              
                    new_note,
                                                                                    
                )
                                                                 
                messagebox.showinfo("Updated", f"Invoice updated successfully.\n{edit_summary}")
                                                                          
                dialog.destroy()
                                                                                       
                self.refresh_recent_invoices()
                                                            
            except Exception as exc:
                                                          
                messagebox.showerror("Error", str(exc))

                                                        
        CTkButton(dialog, text="Submit", fg_color="#f78222", text_color="black",
                                                                                
                  font=("Arial Black", 12), command=save_changes).pack(pady=(5, 15))

                                                                 
    def delete_invoice_dialog(self):
                                                                                
        if not self.selected_invoice:
                                                          
            return

                                                    
        invoice = self.selected_invoice
                                                    
        confirm = messagebox.askyesno("Delete Invoice", f"Delete invoice {invoice[9] or f'RCP-{invoice[0]:05d}'}?")
                                                                                
        if not confirm:
                                                          
            return

                                                                   
        try:
                                                                                   
            self.db.delete_invoice(invoice[0])
                                                             
            messagebox.showinfo("Deleted", "Invoice deleted successfully.")
                                                                                   
            self.refresh_recent_invoices()
                                                                                   
            self.refresh_recent_invoice_details()
                                                                                   
            self.refresh_stat_cards()
                                                                                   
            self.selected_invoice = None
                                                                                   
            self.selected_invoice_id = None
                                                        
        except Exception as exc:
                                                      
            messagebox.showerror("Error", str(exc))

                       

                                                           
def create_invoice_docx(output_path, invoice_data):
                                                
    doc = Document()
                                                
    section = doc.sections[0]
                                                
    section.top_margin = Inches(1.1)
                                                
    section.bottom_margin = Inches(1.0)
                                                
    section.left_margin = Inches(1.1)
                                                
    section.right_margin = Inches(1.1)

                                                
    style = doc.styles["Normal"]
                                                
    style.font.name = "Franklin Gothic Book"
                                                
    style.font.size = Pt(12)

                                                
    title = doc.add_paragraph()
                                                
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                                
    title_run = title.add_run("Invoice")
                                                
    title_run.bold = True
                                                
    title_run.font.size = Pt(22)
                                                
    title_run.font.name = "Franklin Gothic Book"

                                                
    info_table = doc.add_table(rows=2, cols=3)
                                                
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                                                
    info_table.autofit = True

                                                
    payment_date = format_invoice_date(invoice_data["payment_date"])
                                                  
    info_values = [
                                                                  
        ("Receipt no:", invoice_data["receipt_no"]),
                                                                  
        ("Name:", invoice_data["student_name"]),
                                                                  
        ("Class:", invoice_data["student_class"]),
                                                                  
        ("Amount:", f"Rs.{invoice_data['amount']}"),
                                                                  
        ("Date:", payment_date),
                                                                  
        ("", ""),
                                                                        
    ]

                                                        
    for index, (label, value) in enumerate(info_values):
                                                    
        cell = info_table.cell(index // 3, index % 3)
                                                    
        paragraph = cell.paragraphs[0]
                                                    
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                                                    
        label_run = paragraph.add_run(label)
                                                    
        label_run.bold = True
                                                    
        label_run.font.size = Pt(11)
                                                    
        value_run = paragraph.add_run(f" {value}" if label else "")
                                                    
        value_run.bold = True
                                                    
        value_run.font.size = Pt(11)

                                                              
    doc.add_paragraph()
                                                
    bill_table = doc.add_table(rows=4, cols=2)
                                                
    bill_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                                                
    bill_table.style = "Table Grid"
                                                              
    bill_table.columns[0].width = Inches(4.2)
                                                              
    bill_table.columns[1].width = Inches(1.4)

                                                
    headers = ("Description", "Amount")
                                                        
    for col, text in enumerate(headers):
                                                    
        cell = bill_table.cell(0, col)
                                                    
        paragraph = cell.paragraphs[0]
                                                    
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                                    
        run = paragraph.add_run(text)
                                                    
        run.bold = True
                                                    
        run.font.size = Pt(16)

                                                
    description = (
                                                                  
        f"Tuition fee for {invoice_data['month']}\n"
                                                                  
        f"Student ID: {invoice_data['student_id']}\n"
                                                                  
        f"Payment Method: {invoice_data['method']}\n"
                                                                  
        f"Status: {invoice_data['status']}"
                                                                        
    )
                                                              
    bill_table.cell(1, 0).text = description
                                                              
    bill_table.cell(1, 1).text = f"Rs.{invoice_data['amount']}"
                                                              
    bill_table.cell(2, 0).text = ""
                                                              
    bill_table.cell(2, 1).text = ""

                                                
    note_cell = bill_table.cell(3, 0)
                                                
    note_cell.text = ""
                                                
    note_paragraph = note_cell.paragraphs[0]
                                                
    note_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                                
    note_run = note_paragraph.add_run(f"Note: {invoice_data['note']}")
                                                
    note_run.bold = True
                                                
    note_run.font.size = Pt(14)
                                                              
    bill_table.cell(3, 1).text = ""

                                                        
    for row in bill_table.rows:
                                                            
        for cell in row.cells:
                                                        
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

                                                
    total = doc.add_paragraph()
                                                
    total.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                                                
    total_run = total.add_run(f"Total: Rs.{invoice_data['amount']}")
                                                
    total_run.bold = True
                                                
    total_run.font.size = Pt(14)

                                                              
    doc.save(output_path)
                                                  
    return output_path


                                                           
def format_invoice_date(value):
                                                                            
    if hasattr(value, "strftime"):
                                                      
        return value.strftime("%d/%m/%Y")
                                                
    value_text = str(value)
                                                        
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
                                                                   
        try:
                                                          
            return datetime.strptime(value_text, fmt).strftime("%d/%m/%Y")
                                                        
        except ValueError:
                                                                      
            pass
                                                  
    return value_text

                                                             
def generate_invoice_docx(invoice, template_path, output_path):                        
                                                              
    """
    invoice: tuple/row from db.get_recent_invoices() or similar, indexed as:
        invoice[0]  -> invoice id
        invoice[1]  -> student id
        invoice[2]  -> student name
        invoice[3]  -> student class
        invoice[5]  -> amount
        invoice[6]  -> month
        invoice[7]  -> status (P/UN)
        invoice[8]  -> payment date
        invoice[9]  -> receipt no
        invoice[10] -> method
        invoice[11] -> note
    template_path: path to invoice_template.docx
    output_path: where to save the filled .docx
    """
                                                
    doc = DocxTemplate(template_path)
 
                                                
    receipt_no = invoice[9] or f"RCP-{invoice[0]:05d}"
                                                
    amount = invoice[5]
 
                                                        
    context = {
                                                                  
        "RCP": receipt_no,
                                                                  
        "name": invoice[2],
                                                                  
        "class": invoice[3],
                                                                  
        "Amount": f"{amount}",
                                                                  
        "date": invoice[8] or datetime.now().strftime("%d/%m/%Y"),
                                                                  
        "Note": invoice[11] or "-",
                                                                  
        "total": f"{amount}",
                                                                        
    }
 
                                                              
    doc.render(context)
                                                              
    doc.save(output_path)
                                                  
    return output_path
 
 
                                                           
def convert_docx_to_pdf(docx_path, pdf_dir=None):
                                                              
    """
    Converts docx -> pdf.
    Tries docx2pdf first (Windows + MS Word). Falls back to LibreOffice
    headless (soffice) if installed and on PATH.
    Returns the pdf path, or raises if neither method works.
    """
                                                
    pdf_dir = pdf_dir or os.path.dirname(docx_path)
 
                                                                    
                                                               
    try:
                                                                        
        from docx2pdf import convert
                                                                  
        convert(docx_path, pdf_dir)
                                                    
        pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
                                                                                
        if os.path.exists(pdf_path):
                                                          
            return pdf_path
                                                    
    except Exception:
                                                                  
        pass
 
                                                                                  
                                                               
    try:
                                                                  
        subprocess.run(
                                                                      
            ["soffice", "--headless", "--convert-to", "pdf",
                                                                       
             "--outdir", pdf_dir, docx_path],
                                                        
            check=True,
                                                        
            stdout=subprocess.DEVNULL,
                                                        
            stderr=subprocess.DEVNULL,
                                                                            
        )
                                                    
        pdf_path = os.path.join(
                                                                      
            pdf_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf"
                                                                            
        )
                                                                                
        if os.path.exists(pdf_path):
                                                          
            return pdf_path
                                                    
    except Exception as exc:
                                                                  
        raise RuntimeError(
                                                                      
            "PDF conversion failed. Install MS Word (for docx2pdf) or "
                                                                      
            "LibreOffice (for soffice command)."
                                                                            
        ) from exc
 
                                                              
    raise RuntimeError("PDF conversion failed â€” no working converter found.")
